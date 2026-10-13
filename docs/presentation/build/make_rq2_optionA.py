# -*- coding: utf-8 -*-
"""슬라이드 2 대안 A -- 요약 range plot 1개 + 대표 small multiple 2개.
현재 덱(EthosU_RQ_v9.pptx)은 건드리지 않는 별도 파일이다.

데이터: evidence/executability-20260825 (133셀). SSE-300/320 값은
3_3_cycles_by_mac.csv의 canonical 74셀과 일치함을 assert로 확인한다.
"""
import csv, json, math, os, tarfile

exec(open("v8_helpers.py", encoding="utf-8").read())
from pptx.enum.shapes import MSO_CONNECTOR

OUT = "EthosU_RQ_slide2_optionA.pptx"
PAPER = "../../paper"
EXEC_TAR = PAPER + "/evidence/executability-20260825/cells/per_cell_evidence.tar.gz"
CYCLES_CSV = PAPER + "/raw_data_report/tables/3_3_cycles_by_mac.csv"

PALE = RGBColor(0xF7, 0xF9, 0xFC)
LINEC = RGBColor(0xCB, 0xD6, 0xE3)
MUTED = RGBColor(0x8A, 0x97, 0xA6)
EYE = RGBColor(0xAE, 0xC2, 0xDE)
SPAN = RGBColor(0xC4, 0xCF, 0xDC)

MODELS = [
    ("wav2letter_pruned_int8", "Wav2Letter", RGBColor(0x13, 0x36, 0x6B)),
    ("mobilenet_v2_1.0_224_INT8", "MobileNetV2", RGBColor(0x4A, 0x7A, 0xB8)),
    ("yolo-fastest_192_face_v4", "YOLO-Fastest", RGBColor(0x83, 0xAB, 0xD6)),
    ("ad_medium_int8", "AD MedNet", RGBColor(0x2E, 0x7D, 0x4F)),
    ("vww4_128_128_INT8", "VWW4", RGBColor(0x8C, 0x9A, 0xA8)),
    ("kws_micronet_m", "KWS MicroNet", RGBColor(0xC2, 0x8A, 0x1E)),
    ("rnnoise_INT8", "RNNoise", RGBColor(0xA4, 0x0F, 0x16)),
]
ENDS = {"wav2letter_pruned_int8", "rnnoise_INT8"}
CFG = [("Corstone-300 · U55", "SSE-300", "ethos-u55"),
       ("Corstone-300 · U65", "SSE-300", "ethos-u65"),
       ("Corstone-310 · U55", "SSE-310", "ethos-u55"),
       ("Corstone-310 · U65", "SSE-310", "ethos-u65"),
       ("Corstone-315 · U65", "SSE-315", "ethos-u65"),
       ("Corstone-320 · U85", "SSE-320", "ethos-u85")]
SHOW = [CFG[0], CFG[5]]           # 대표 패널 두 개

CYC = {}
_tf = tarfile.open(EXEC_TAR)
for _m in _tf.getmembers():
    _n = os.path.basename(_m.name)
    if not _n.endswith(".json") or _n.startswith("_"):
        continue
    _d = json.load(_tf.extractfile(_m))
    _wl, _plat, _nm = _d["cell_id"].split("__")
    _npu, _mac = _nm.rsplit("-", 1)
    for _a in _d["attempts"]:
        _r = _a.get("run") or {}
        if _r.get("status") == "SUCCESS" and _r.get("npu_total_cycles") is not None:
            CYC[(_plat, _npu, _wl, int(_mac))] = _r["npu_total_cycles"]
_k = 0
for _r in csv.DictReader(open(CYCLES_CSV, encoding="utf-8")):
    if _r["cycles"]:
        assert CYC[(_r["platform"], _r["npu"], _r["workload"], int(_r["mac"]))] == int(_r["cycles"])
        _k += 1
assert _k == 74

VAL = {m: [CYC[(p, u, m, 256)] for _l, p, u in CFG] for m, _n, _c in MODELS}
# 순위가 여섯 구성에서 모두 같은지 (화면 문장의 근거)
_ref = None
for _l, _p, _u in CFG:
    _o = sorted(VAL, key=lambda w: CYC[(_p, _u, w, 256)])
    _ref = _o if _ref is None else _ref
    assert _o == _ref, _l
SPREAD = [max(v) / min(v) for v in VAL.values()]
GAPS = [max(CYC[(p, u, m, 256)] for m, _n, _c in MODELS)
        / min(CYC[(p, u, m, 256)] for m, _n, _c in MODELS) for _l, p, u in CFG]


def med(v):
    t = sorted(v)
    return (t[2] + t[3]) / 2


def blank():
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s


def conn(s, x1, y1, x2, y2, color=LINEC, w=0.75, dash=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                               Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": dash}))
    return c


def dot(s, cx, cy, d, color, ring=None):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2),
                           Inches(d), Inches(d))
    o.fill.solid(); o.fill.fore_color.rgb = color
    if ring is None:
        o.line.fill.background()
    else:
        o.line.color.rgb = ring; o.line.width = Pt(1.0)
    o.shadow.inherit = False
    return o


LO, HI = 4.13, 7.37


def cxl(v, x0, w):
    return x0 + w * (math.log10(v) - LO) / (HI - LO)


for _sid in list(prs.slides._sldIdLst):
    prs.part.drop_rel(_sid.rId)
    prs.slides._sldIdLst.remove(_sid)

s = blank()
tb(s, 0.48, 0.26, 8.90, 0.66, "2. 모델별 실행 비용은 어떻게 다른가?", size=26, bold=True,
   color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 9.45, 0.42, 3.40, 0.30, "RQ1 · 실행 비용", size=11.5, color=EYE, align=PP_ALIGN.RIGHT)
rect(s, 0.50, 1.24, 12.33, 0.84, LIGHT)
rect(s, 0.50, 1.24, 0.055, 0.84, NAVY)
tb(s, 0.74, 1.30, 11.90, 0.72,
   ["여섯 구성 모두에서 모델별 비용 순위가 동일하며,",
    "RNNoise와 Wav2Letter가 실행 비용 범위의 양끝에 놓인다."],
   size=15, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE, spacing=1.20)
tb(s, 0.52, 6.56, 12.30, 0.28,
   "256 MAC · TOTAL cycles · 공통 log scale · 왼쪽은 여섯 구성의 최소–최대와 각 값, 굵은 점은 중앙값   |   "
   "모델 간 차이 146–772배 > 구성 간 변동 1.1–5.5배 · 나머지 네 구성도 같은 순서", size=9.5, color=MUTED)

# ------------------------------------------------ 왼쪽: 모델별 비용 범위 요약
tb(s, 0.50, 2.26, 7.30, 0.26, "모델별 비용 범위 — 여섯 구성", size=11.5, bold=True, color=NAVY)
LIX, LIW = 2.16, 5.62
ROW0, ROWH = 2.74, 0.50
for g in (1e5, 1e6, 1e7):
    gx = cxl(g, LIX, LIW)
    conn(s, gx, ROW0, gx, ROW0 + 7 * ROWH, LINEC, 0.75, "sysDot")
    tb(s, gx - 0.40, ROW0 + 7 * ROWH + 0.04, 0.80, 0.22,
       {1e5: "100K", 1e6: "1M", 1e7: "10M"}[g], size=9, color=MUTED, align=PP_ALIGN.CENTER)
for i, (mid, name, colr) in enumerate(MODELS):
    yy = ROW0 + i * ROWH
    cy = yy + ROWH / 2
    end = mid in ENDS
    tb(s, 0.50, yy, 1.54, ROWH, name, size=11.5 if end else 10.5, bold=end,
       color=colr if end else GRAY, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    v = VAL[mid]
    conn(s, cxl(min(v), LIX, LIW), cy, cxl(max(v), LIX, LIW), cy, SPAN, 3.4)
    for x in v:
        dot(s, cxl(x, LIX, LIW), cy, 0.095, RGBColor(0x9F, 0xAE, 0xC0), WHITE)
    dot(s, cxl(med(v), LIX, LIW), cy, 0.175, colr, WHITE)

# ------------------------------------------------ 오른쪽: 대표 구성 두 개
tb(s, 8.15, 2.26, 4.68, 0.26, "대표 구성 — 행 순서와 색은 왼쪽과 같다", size=11.5,
   bold=True, color=NAVY)
RX, RW = 8.15, 4.68
for pi, (label, plat, npu) in enumerate(SHOW):
    hy = 2.58 + pi * 1.88
    rect(s, RX, hy, RW, 0.26, NAVY)
    tb(s, RX, hy, RW, 0.26, label, size=9.5, bold=True, color=WHITE,
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    by, bh = hy + 0.26, 1.42
    rect(s, RX, by, RW, bh, PALE)
    ix, iw = RX + 0.20, RW - 0.40
    for g in (1e5, 1e6, 1e7):
        gx = cxl(g, ix, iw)
        conn(s, gx, by + 0.05, gx, by + bh - 0.05, LINEC, 0.75, "sysDot")
    for i, (mid, _n, colr) in enumerate(MODELS):
        v = CYC[(plat, npu, mid, 256)]
        dot(s, cxl(v, ix, iw), by + (i + 0.5) * (bh / 7),
            0.135 if mid in ENDS else 0.10, colr, WHITE)
    if pi == 1:
        for g in (1e5, 1e6, 1e7):
            tb(s, cxl(g, ix, iw) - 0.40, by + bh + 0.04, 0.80, 0.22,
               {1e5: "100K", 1e6: "1M", 1e7: "10M"}[g], size=9, color=MUTED,
               align=PP_ALIGN.CENTER)
prs.save(OUT)
print("wrote", OUT)
print("구성 간 변동 %.2f-%.2f배 / 구성 안 모델 간 차이 %.0f-%.0f배"
      % (min(SPREAD), max(SPREAD), min(GAPS), max(GAPS)))
