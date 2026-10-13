# -*- coding: utf-8 -*-
"""6×6 과 36×1 의 block 경로 대조 (3×3 Conv · U85-256 · 기본 TA).
값은 h3_blocks.csv 에서 빌드 시점에 읽는다.
"""
import csv
exec(open("v8_helpers.py", encoding="utf-8").read())
from pptx.enum.shapes import MSO_CONNECTOR

OUT = "EthosU_RQ_blockchain.pptx"
BLK = "../../paper/raw_data_report/amendments/h13/h3_blocks.csv"
MUTED = RGBColor(0x8A, 0x97, 0xA6)
EYE = RGBColor(0xAE, 0xC2, 0xDE)
CARD = RGBColor(0xF4, 0xF6, 0xF9)
CHIP = RGBColor(0xF1, 0xF4, 0xF8)
PILL = RGBColor(0xEE, 0xF2, 0xF7)
EDGE = RGBColor(0xD8, 0xE0, 0xEA)
STEEL = RGBColor(0x4A, 0x7A, 0xB8)

R = {}
for r in csv.DictReader(open(BLK, encoding="utf-8")):
    if r["op"] == "conv3x3" and r["area"] == "36" and r["cond"] == "base" and r["mac"] == "256":
        R[(int(r["H"]), int(r["W"]))] = r
A, B = R[(6, 6)], R[(36, 1)]
assert A["ublock_HWC"] == B["ublock_HWC"], "ublock 이 같아야 대조가 성립"


def blank():
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s


def card(s, x, y, w, h, fill=CARD, edge=EDGE):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = fill
    r.line.color.rgb = edge; r.line.width = Pt(0.9); r.shadow.inherit = False
    r.adjustments[0] = 0.06
    return r


ARROW = RGBColor(0xA8, 0xB4, 0xC2)


def right(s, x, y, w, h, label=None):
    """가느다란 가로 화살표 — 채운 도형이 아니라 글자."""
    tb(s, x, y, w, h, "→", size=15, color=ARROW, align=PP_ALIGN.CENTER,
       anchor=MSO_ANCHOR.MIDDLE)
    if label:
        tb(s, x - 0.30, y + h * 0.5, w + 0.60, 0.20, label, size=8.5, bold=True,
           color=ARROW, align=PP_ALIGN.CENTER)


def step(s, x, y, w, h, tag, value, note, notecol):
    card(s, x, y, w, h, CARD, CARD)
    tb(s, x + 0.16, y + 0.10, w - 0.32, 0.24, tag, size=10, bold=True, color=GRAY)
    tb(s, x + 0.16, y + 0.40, w - 0.32, 0.38, value, size=16, bold=True, color=TEXT,
       anchor=MSO_ANCHOR.MIDDLE)
    tb(s, x + 0.16, y + 0.84, w - 0.32, 0.44, note, size=9.5, color=notecol, spacing=1.18)


def fmt(n):
    return "{:,}".format(int(n))


for _sid in list(prs.slides._sldIdLst):
    prs.part.drop_rel(_sid.rId)
    prs.slides._sldIdLst.remove(_sid)

s = blank()
tb(s, 0.48, 0.26, 8.90, 0.66, "같은 면적, 다른 형상 — block 경로가 갈린다", size=26, bold=True,
   color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 9.45, 0.42, 3.40, 0.30, "3×3 Conv · U85-256", size=11.5, color=EYE, align=PP_ALIGN.RIGHT)
tb(s, 0.52, 1.20, 12.30, 0.40,
   "공간 원소는 둘 다 36개다. 논리 폭이 ublock 폭의 배수가 아니면 block 기하가 달라지고, "
   "그것이 블록 수와 활용률을 가른다.",
   size=13.5, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)

X0, TW, RH = 0.50, 12.33, 2.06
XA, WA = 0.70, 2.16          # ① 형상
XB, WB = 3.18, 2.32          # ② OFM block
XC, WC = 6.34, 2.42          # ③ IFM block
XD, WD = 9.10, 2.40          # 지표 칩 3개
XE, WE = 11.62, 1.14         # PMU
for ri, (row, name, colr) in enumerate([(A, "6×6", STEEL), (B, "36×1", RED)]):
    y = 1.70 + ri * (RH + 0.18)
    card(s, X0, y, TW, RH, WHITE, EDGE)
    H, W = int(row["H"]), int(row["W"])
    fit = W > 1
    sy = y + 0.50                                   # 단계 카드 상단
    SH = 1.36

    tb(s, XA, y + 0.12, WA, 0.30, "%s feature map" % name, size=14, bold=True, color=colr,
       anchor=MSO_ANCHOR.MIDDLE)
    PIT, CELL = 0.092, 0.082
    gx, gy = XA + 0.04, sy + 0.06
    if fit:
        for rr in range(H):
            for cc in range(W):
                rect(s, gx + cc * PIT, gy + rr * PIT, CELL, CELL, colr)
        gw = W * PIT
    else:
        rect(s, gx + 0.16, gy - 0.02, 0.15, 1.18, colr)
        gw = 0.40
    tb(s, gx + gw + 0.26, sy + 0.28, 1.60, 0.32, "%d × %d × C" % (H, W), size=14, bold=True,
       color=TEXT, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, gx + gw + 0.26, sy + 0.64, 1.60, 0.46,
       "공간 원소 36개" if fit else "공간 원소는\n똑같이 36개", size=9.5, color=GRAY, spacing=1.18)

    right(s, XA + WA + 0.02, sy, 0.28, SH)
    step(s, XB, sy, WB, SH, "② Vela OFM block", row["ofm_block_HWC"].replace("x", " × "),
         "공간 방향이 논리 OFM 과 맞음" if fit else "논리 폭 W=1 인데\nblock 은 W=2 로 잡힘",
         GRAY if fit else RED)
    right(s, XB + WB + 0.04, sy, 0.50, SH, "3×3 halo")
    _i = [int(v) for v in row["ifm_block"].split("x")]
    step(s, XC, sy, WC, SH, "③ IFM block", row["ifm_block"].replace("x", " × "),
         "입력 원소 %s개" % fmt(_i[0] * _i[1] * _i[2] * _i[3]), GRAY if fit else RED)
    right(s, XC + WC + 0.04, sy, 0.28, SH)

    GAP = 0.08
    cwid = (WD - GAP * 2) / 3
    for k, (lab, val) in enumerate([("uBlock", row["ublock_HWC"].replace("x", "×")),
                                    ("blocks", row["n_blocks"]),
                                    ("util", row["vela_util_pct"] + "%")]):
        bx = XD + k * (cwid + GAP)
        card(s, bx, sy + 0.34, cwid, 0.70, CHIP, CHIP)
        tb(s, bx, sy + 0.38, cwid, 0.20, lab, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
        tb(s, bx, sy + 0.56, cwid, 0.28, val, size=13.5, bold=True,
           color=colr if (k and not fit) else NAVY, align=PP_ALIGN.CENTER)

    card(s, XE, sy + 0.34, WE, 0.70, PILL, PILL)
    tb(s, XE, sy + 0.38, WE, 0.20, "PMU cycles", size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
    tb(s, XE, sy + 0.56, WE, 0.28, fmt(row["pmu_total"]), size=14.5, bold=True,
       color=NAVY if fit else RED, align=PP_ALIGN.CENTER)

_d = int(B["pmu_total"]) - int(A["pmu_total"])
tb(s, 0.52, 6.12, 12.30, 0.34,
   "같은 ublock(%s), 같은 면적 36 — 논리 폭만 달라 block 기하가 갈리고 블록 수가 %s배, 비용이 +%s cycles (%.0f%%)"
   % (A["ublock_HWC"].replace("x", "×"), int(B["n_blocks"]) // int(A["n_blocks"]), fmt(_d),
      100 * _d / int(A["pmu_total"])),
   size=12.5, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 0.52, 6.50, 12.30, 0.20,
   "출처 h3_blocks.csv (Vela --verbose-schedule) · 기본 TA · 두 형상은 가중치·채널·커널이 같고 H×W 배치만 다르다",
   size=8, color=MUTED)

prs.save(OUT)
print("wrote", OUT)
