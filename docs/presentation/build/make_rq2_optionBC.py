# -*- coding: utf-8 -*-
"""슬라이드 2 대안 B(히트맵)·C(rank plot) -- optionA와 비교용 별도 파일.
데이터·검증은 make_rq2_optionA.py와 동일하다.
"""
import csv, json, math, os, tarfile

exec(open("v8_helpers.py", encoding="utf-8").read())
from pptx.enum.shapes import MSO_CONNECTOR

OUT = "EthosU_RQ_slide2_optionBC.pptx"
PAPER = "../../paper"
EXEC_TAR = PAPER + "/evidence/executability-20260825/cells/per_cell_evidence.tar.gz"
CYCLES_CSV = PAPER + "/raw_data_report/tables/3_3_cycles_by_mac.csv"

PALE = RGBColor(0xF7, 0xF9, 0xFC)
LINEC = RGBColor(0xCB, 0xD6, 0xE3)
MUTED = RGBColor(0x8A, 0x97, 0xA6)
EYE = RGBColor(0xAE, 0xC2, 0xDE)

MODELS = [
    ("wav2letter_pruned_int8", "Wav2Letter", RGBColor(0x13, 0x36, 0x6B)),
    ("mobilenet_v2_1.0_224_INT8", "MobileNetV2", RGBColor(0x4A, 0x7A, 0xB8)),
    ("yolo-fastest_192_face_v4", "YOLO-Fastest", RGBColor(0x83, 0xAB, 0xD6)),
    ("ad_medium_int8", "AD MedNet", RGBColor(0x2E, 0x7D, 0x4F)),
    ("vww4_128_128_INT8", "VWW4", RGBColor(0x8C, 0x9A, 0xA8)),
    ("kws_micronet_m", "KWS MicroNet", RGBColor(0xC2, 0x8A, 0x1E)),
    ("rnnoise_INT8", "RNNoise", RGBColor(0xA4, 0x0F, 0x16)),
]
ENDS = {"wav2letter_pruned_int8", "rnnoise_INT8"}
CFG = [("Corstone-300", "U55", "SSE-300", "ethos-u55"),
       ("Corstone-300", "U65", "SSE-300", "ethos-u65"),
       ("Corstone-310", "U55", "SSE-310", "ethos-u55"),
       ("Corstone-310", "U65", "SSE-310", "ethos-u65"),
       ("Corstone-315", "U65", "SSE-315", "ethos-u65"),
       ("Corstone-320", "U85", "SSE-320", "ethos-u85")]

CYC = {}
_tf = tarfile.open(EXEC_TAR)
for _m in _tf.getmembers():
    _n = os.path.basename(_m.name)
    if not _n.endswith(".json") or _n.startswith("_"):
        continue
    _d = json.load(_tf.extractfile(_m))
    _wl, _plat, _nm = _d["cell_id"].split("__")
    _npu, _mac = _nm.rsplit("-", 1)
    for _a in _d["attempts"]:
        _r = _a.get("run") or {}
        if _r.get("status") == "SUCCESS" and _r.get("npu_total_cycles") is not None:
            CYC[(_plat, _npu, _wl, int(_mac))] = _r["npu_total_cycles"]
_k = 0
for _r in csv.DictReader(open(CYCLES_CSV, encoding="utf-8")):
    if _r["cycles"]:
        assert CYC[(_r["platform"], _r["npu"], _r["workload"], int(_r["mac"]))] == int(_r["cycles"])
        _k += 1
assert _k == 74

V = {(m, i): CYC[(p, u, m, 256)] for m, _n, _c in MODELS
     for i, (_a, _b, p, u) in enumerate(CFG)}
VMIN, VMAX = min(V.values()), max(V.values())
RANK = {}
for i in range(len(CFG)):
    order = sorted((m for m, _n, _c in MODELS), key=lambda m: -V[(m, i)])
    for r, m in enumerate(order, 1):
        RANK[(m, i)] = r
FLIPS = sum(1 for m, _n, _c in MODELS for i in range(1, len(CFG))
            if RANK[(m, i)] != RANK[(m, 0)])
assert FLIPS == 0, FLIPS


def blank():
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s


def head(s, eyebrow, title, answer, caption=None):
    tb(s, 0.48, 0.26, 8.90, 0.66, title, size=26, bold=True, color=WHITE,
       anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 9.45, 0.42, 3.40, 0.30, eyebrow, size=11.5, color=EYE, align=PP_ALIGN.RIGHT)
    rect(s, 0.50, 1.24, 12.33, 0.84, LIGHT)
    rect(s, 0.50, 1.24, 0.055, 0.84, NAVY)
    tb(s, 0.74, 1.30, 11.90, 0.72, answer, size=15, bold=True, color=NAVY,
       anchor=MSO_ANCHOR.MIDDLE, spacing=1.20)
    if caption:
        tb(s, 0.52, 6.56, 12.30, 0.28, caption, size=9.5, color=MUTED)


def conn(s, x1, y1, x2, y2, color=LINEC, w=0.75, dash=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                               Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": dash}))
    return c


def dot(s, cx, cy, d, color, ring=None):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2),
                           Inches(d), Inches(d))
    o.fill.solid(); o.fill.fore_color.rgb = color
    if ring is None:
        o.line.fill.background()
    else:
        o.line.color.rgb = ring; o.line.width = Pt(1.0)
    o.shadow.inherit = False
    return o


def ramp(t):
    a, b = (0xEC, 0xF2, 0xF9), (0x10, 0x2C, 0x5C)
    return RGBColor(*(int(a[i] + (b[i] - a[i]) * t) for i in range(3)))


def shade(v):
    return (math.log10(v) - math.log10(VMIN)) / (math.log10(VMAX) - math.log10(VMIN))


def fmt(n):
    return "{:,}".format(n)


for _sid in list(prs.slides._sldIdLst):
    prs.part.drop_rel(_sid.rId)
    prs.slides._sldIdLst.remove(_sid)

# =============================================================== 대안 B · 히트맵
s = blank()
head(s, "대안 B · 히트맵", "2. 모델별 실행 비용은 어떻게 다른가?",
     ["여섯 구성 모두에서 모델별 비용 순위가 동일하며,",
      "RNNoise와 Wav2Letter가 실행 비용 범위의 양끝에 놓인다."],
     "256 MAC · TOTAL cycles · 칸의 명도는 log(TOTAL cycles) · 진할수록 비싸다   |   "
     "모델 간 차이 146–772배 > 구성 간 변동 1.1–5.5배")

LBLW, CW, HX = 1.84, 1.72, 0.62
HY, HH, RH = 2.34, 0.46, 0.482
for j, (plat, npu, _p, _u) in enumerate(CFG):
    cx = HX + LBLW + j * CW
    rect(s, cx, HY, CW - 0.02, HH, NAVY)
    tb(s, cx, HY, CW - 0.02, HH, [plat, npu], size=9, bold=True, color=WHITE,
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.10)
for i, (mid, name, colr) in enumerate(MODELS):
    yy = HY + HH + i * RH
    end = mid in ENDS
    tb(s, HX, yy, LBLW - 0.14, RH, name, size=11.5 if end else 10.5, bold=end,
       color=colr if end else GRAY, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    for j in range(len(CFG)):
        cx = HX + LBLW + j * CW
        t = shade(V[(mid, j)])
        rect(s, cx, yy, CW - 0.02, RH - 0.02, ramp(t))
        tb(s, cx, yy, CW - 0.02, RH - 0.02, fmt(V[(mid, j)]), size=10, bold=end,
           color=WHITE if t > 0.52 else NAVY, align=PP_ALIGN.CENTER,
           anchor=MSO_ANCHOR.MIDDLE)
GY = HY + HH + 7 * RH + 0.12
tb(s, HX, GY, 1.70, 0.24, "옅음 = 저비용", size=9, color=MUTED, align=PP_ALIGN.RIGHT)
for k in range(24):
    rect(s, HX + 1.80 + k * 0.118, GY + 0.04, 0.12, 0.16, ramp(k / 23.0))
tb(s, HX + 4.68, GY, 1.70, 0.24, "진함 = 고비용", size=9, color=MUTED)

# =============================================================== 대안 C · rank plot
s = blank()
head(s, "대안 C · 순위 + 실제 값", "2. 모델별 실행 비용은 어떻게 다른가?",
     ["여섯 구성에서 모델 순위가 한 번도 바뀌지 않는다 (순위 변동 0회).",
      "RNNoise와 Wav2Letter가 실행 비용 범위의 양끝에 놓인다."],
     "왼쪽: 구성별 비용 순위(1 = 가장 비쌈) · 오른쪽: 대표 구성의 실제 값, 256 MAC · TOTAL cycles · 로그축   |   "
     "모델 간 차이 146–772배 > 구성 간 변동 1.1–5.5배")

tb(s, 0.50, 2.26, 7.00, 0.26, "구성별 비용 순위", size=11.5, bold=True, color=NAVY)
RX0, RXW = 2.34, 4.62
RY0, RYH = 2.86, 0.53
for j, (plat, npu, _p, _u) in enumerate(CFG):
    gx = RX0 + RXW * j / 5.0
    conn(s, gx, RY0 - 0.06, gx, RY0 + 6 * RYH + 0.10, LINEC, 0.75, "sysDot")
    tb(s, gx - 0.48, RY0 + 6 * RYH + 0.14, 0.96, 0.40, [plat.replace("Corstone-", ""), npu],
       size=8.5, color=MUTED, align=PP_ALIGN.CENTER, spacing=1.12)
for i, (mid, name, colr) in enumerate(MODELS):
    end = mid in ENDS
    ys = [RY0 + (RANK[(mid, j)] - 1) * RYH for j in range(len(CFG))]
    xs = [RX0 + RXW * j / 5.0 for j in range(len(CFG))]
    for a in range(len(CFG) - 1):
        conn(s, xs[a], ys[a], xs[a + 1], ys[a + 1], colr, 2.6 if end else 1.6)
    for x, y in zip(xs, ys):
        dot(s, x, y, 0.125 if end else 0.095, colr, WHITE)
    tb(s, 0.50, ys[0] - 0.13, 1.72, 0.26, name, size=11 if end else 10, bold=end,
       color=colr if end else GRAY, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, xs[-1] + 0.14, ys[-1] - 0.13, 0.40, 0.26, str(RANK[(mid, 0)]), size=10,
       bold=end, color=colr if end else MUTED, anchor=MSO_ANCHOR.MIDDLE)

tb(s, 7.72, 2.26, 5.11, 0.26, "대표 구성 · Corstone-300 · U55", size=11.5, bold=True,
   color=NAVY)
bx, by, bw, bh = 7.72, 2.60, 5.11, 3.52
rect(s, bx, by, bw, bh, PALE)
LO, HI = 4.13, 7.37
ix, iw = bx + 0.24, bw - 0.48
for g in (1e5, 1e6, 1e7):
    gx = ix + iw * (math.log10(g) - LO) / (HI - LO)
    conn(s, gx, by + 0.06, gx, by + bh - 0.32, LINEC, 0.75, "sysDot")
    tb(s, gx - 0.40, by + bh - 0.30, 0.80, 0.22, {1e5: "100K", 1e6: "1M", 1e7: "10M"}[g],
       size=9, color=MUTED, align=PP_ALIGN.CENTER)
for i, (mid, name, colr) in enumerate(MODELS):
    v = V[(mid, 0)]
    gx = ix + iw * (math.log10(v) - LO) / (HI - LO)
    cy = by + 0.20 + (i + 0.5) * ((bh - 0.56) / 7)
    dot(s, gx, cy, 0.145 if mid in ENDS else 0.105, colr, WHITE)

prs.save(OUT)
print("wrote", OUT, "| 순위 변동 횟수:", FLIPS)
