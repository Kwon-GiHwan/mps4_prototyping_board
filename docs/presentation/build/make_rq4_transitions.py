# -*- coding: utf-8 -*-
"""초안 -- v9 4장(MAC 확장) 대안.
왼쪽 65% 인접 MAC 전이별 speedup heatmap (7모델 × 13전이),
오른쪽 35% RNNoise 프로파일 교차 그래프 (네 점, 고정선 2 + 기본 조합 점선 1).
별도 파일이며 EthosU_RQ_v9.pptx 는 건드리지 않는다.
"""
import csv, json, math, os, tarfile

exec(open("v8_helpers.py", encoding="utf-8").read())
from pptx.enum.shapes import MSO_CONNECTOR

OUT = "EthosU_RQ_slide4_transitions.pptx"
PAPER = "../../paper"
EXEC_TAR = PAPER + "/evidence/executability-20260825/cells/per_cell_evidence.tar.gz"
CYCLES_CSV = PAPER + "/raw_data_report/tables/3_3_cycles_by_mac.csv"

PALE = RGBColor(0xF7, 0xF9, 0xFC)
LINEC = RGBColor(0xCB, 0xD6, 0xE3)
MUTED = RGBColor(0x8A, 0x97, 0xA6)
EYE = RGBColor(0xAE, 0xC2, 0xDE)
GREEN = RGBColor(0x2E, 0x7D, 0x4F)
STEEL = RGBColor(0x4A, 0x7A, 0xB8)
DEAD = RGBColor(0xF1, 0xF3, 0xF5)

MODELS = [("wav2letter_pruned_int8", "Wav2Letter"),
          ("mobilenet_v2_1.0_224_INT8", "MobileNetV2"),
          ("yolo-fastest_192_face_v4", "YOLO-Fastest"),
          ("ad_medium_int8", "AD MedNet"),
          ("vww4_128_128_INT8", "VWW4"),
          ("kws_micronet_m", "KWS MicroNet"),
          ("rnnoise_INT8", "RNNoise")]
HERO = "rnnoise_INT8"
GROUPS = [("Corstone-300 · U55", "SSE-300", "ethos-u55", [32, 64, 128, 256]),
          ("Corstone-300 · U65", "SSE-300", "ethos-u65", [256, 512]),
          ("Corstone-310 · U55", "SSE-310", "ethos-u55", [32, 64, 128, 256]),
          ("Corstone-310 · U65", "SSE-310", "ethos-u65", [256, 512]),
          ("Corstone-315 · U65", "SSE-315", "ethos-u65", [256, 512]),
          ("Corstone-320 · U85", "SSE-320", "ethos-u85", [128, 256, 512, 1024, 2048])]

CYC = {}
_tf = tarfile.open(EXEC_TAR)
for _m in _tf.getmembers():
    _n = os.path.basename(_m.name)
    if not _n.endswith(".json") or _n.startswith("_"):
        continue
    _d = json.load(_tf.extractfile(_m))
    _wl, _plat, _nm = _d["cell_id"].split("__")
    _npu, _mac = _nm.rsplit("-", 1)
    for _a in _d["attempts"]:
        _r = _a.get("run") or {}
        if _r.get("status") == "SUCCESS" and _r.get("npu_total_cycles") is not None:
            CYC[(_plat, _npu, _wl, int(_mac))] = _r["npu_total_cycles"]
_k = 0
for _r in csv.DictReader(open(CYCLES_CSV, encoding="utf-8")):
    if _r["cycles"]:
        assert CYC[(_r["platform"], _r["npu"], _r["workload"], int(_r["mac"]))] == int(_r["cycles"])
        _k += 1
assert _k == 74

TRANS = []                                   # (group_index, plat, npu, a, b)
for gi, (_l, p, u, macs) in enumerate(GROUPS):
    for a, b in zip(macs, macs[1:]):
        TRANS.append((gi, p, u, a, b))
assert len(TRANS) == 13, len(TRANS)


def speedup(mid, t):
    _gi, p, u, a, b = t
    x, y = CYC.get((p, u, mid, a)), CYC.get((p, u, mid, b))
    return x / y if (x and y) else None


SP = {(m, i): speedup(m, t) for m, _n in MODELS for i, t in enumerate(TRANS)}
REV = [(m, i) for (m, i), s in SP.items() if s is not None and s < 1.0]
assert sorted(m for m, _i in REV) == [HERO, HERO], REV   # 역전은 RNNoise 2건뿐
HERO_COL = next(i for i, t in enumerate(TRANS)
                if t[1] == "SSE-320" and t[3] == 256 and t[4] == 512)


def blank():
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s


def conn(s, x1, y1, x2, y2, color=LINEC, w=0.75, dash=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                               Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": dash}))
    return c


def dot(s, cx, cy, d, color, ring=None):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2),
                           Inches(d), Inches(d))
    o.fill.solid(); o.fill.fore_color.rgb = color
    if ring is None:
        o.line.fill.background()
    else:
        o.line.color.rgb = ring; o.line.width = Pt(1.0)
    o.shadow.inherit = False
    return o


def outline(s, x, y, w, h, color, pt=2.0):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.background(); r.line.color.rgb = color; r.line.width = Pt(pt)
    r.shadow.inherit = False
    return r


def lerp(a, b, t):
    return RGBColor(*(int(a[i] + (b[i] - a[i]) * t) for i in range(3)))


NEUTRAL = (0xF0, 0xF4, 0xF8)
GOOD = (0x13, 0x36, 0x6B)
BAD = (0xA4, 0x0F, 0x16)


def heat(s_):
    """2x = 진한 개선색, 1x = 중립, 1x 미만 = 역전 강조색."""
    if s_ is None:
        return DEAD, MUTED
    if s_ >= 1.0:
        t = min((s_ - 1.0) / 1.0, 1.0)
        return lerp(NEUTRAL, GOOD, t), (WHITE if t > 0.45 else RGBColor(0x22, 0x36, 0x4B))
    t = min((1.0 - s_) / 0.40, 1.0)
    return lerp((0xF9, 0xE8, 0xE8), BAD, t), (WHITE if t > 0.45 else RGBColor(0xA4, 0x0F, 0x16))


for _sid in list(prs.slides._sldIdLst):
    prs.part.drop_rel(_sid.rId)
    prs.slides._sldIdLst.remove(_sid)

s = blank()
tb(s, 0.48, 0.26, 8.90, 0.66, "4. MAC 증가의 이득은 언제 줄거나 역전되는가?", size=26,
   bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 9.45, 0.42, 3.40, 0.30, "RQ1 · MAC 확장", size=11.5, color=EYE, align=PP_ALIGN.RIGHT)
rect(s, 0.50, 1.24, 12.33, 0.84, LIGHT)
rect(s, 0.50, 1.24, 0.055, 0.84, NAVY)
tb(s, 0.74, 1.30, 11.90, 0.72,
   ["확장 이득은 모델·전이마다 다르며,",
    "RNNoise의 첫 역전은 메모리 프로파일 변경과 연결된다."],
   size=15, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE, spacing=1.20)
tb(s, 0.52, 6.52, 12.30, 0.36,
   ["왼쪽 Speedup = C(M) / C(2M) · 2× = MAC 두 배에 비례한 개선 · 1× = 변화 없음 · 1× 미만 = 역전 · "
    "— = 두 측정값 중 하나가 없는 셀   |   오른쪽 RNNoise · U85 · 256 MAC → 512 MAC · TOTAL cycles"],
   size=9, color=MUTED)

# ---------------------------------------------- 왼쪽 65%: 인접 전이 heatmap
LX, LBW = 0.50, 1.40
HX = LX + LBW + 0.04
CW = 0.518
GY, GH = 2.34, 0.32          # 플랫폼 그룹 머리
TY, TH = 2.68, 0.30          # 전이 라벨
RY, RH = 3.02, 0.44          # 데이터 행

cx = HX
for gi, (glbl, _p, _u, macs) in enumerate(GROUPS):
    n = len(macs) - 1
    rect(s, cx, GY, n * CW - 0.02, GH, NAVY)
    tb(s, cx, GY, n * CW - 0.02, GH, glbl.replace("Corstone-", ""), size=8.5, bold=True,
       color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    cx += n * CW
def _m(v):
    return "%dk" % (v // 1024) if v >= 1024 else str(v)


for i, (_gi, _p, _u, a, b) in enumerate(TRANS):
    tb(s, HX + i * CW, TY, CW - 0.02, TH, "%s→%s" % (_m(a), _m(b)), size=7, color=GRAY,
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

for r, (mid, name) in enumerate(MODELS):
    yy = RY + r * RH
    hero = mid == HERO
    tb(s, LX, yy, LBW, RH, name, size=10.5, bold=hero,
       color=RED if hero else TEXT, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    for i in range(len(TRANS)):
        v = SP[(mid, i)]
        fill, txt = heat(v)
        rect(s, HX + i * CW, yy, CW - 0.02, RH - 0.02, fill)
        show = hero or (v is not None and v < 1.0)
        label = "—" if v is None else ("%.2f" % v if show else "")
        if label:
            tb(s, HX + i * CW, yy, CW - 0.02, RH - 0.02, label, size=8, bold=show,
               color=txt, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
outline(s, HX + HERO_COL * CW - 0.02, RY + 6 * RH - 0.02, CW + 0.02, RH + 0.02, RED, 2.25)

# 색 범례
lgx = HX
tb(s, LX, 6.16, LBW, 0.24, "Speedup", size=8.5, bold=True, color=NAVY,
   align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
for k, (lo, hi, lab) in enumerate([(0.60, 1.0, None), (1.0, 2.0, None)]):
    for j in range(16):
        t = j / 15.0
        v = lo + (hi - lo) * t
        f, _c = heat(v)
        rect(s, lgx + (k * 16 + j) * 0.072, 6.20, 0.075, 0.16, f)
for gx, lab in ((0, "0.6×"), (16 * 0.072, "1×"), (32 * 0.072, "2×")):
    tb(s, HX + gx - 0.30, 6.37, 0.60, 0.22, lab, size=7.5, color=MUTED,
       align=PP_ALIGN.CENTER)
tb(s, HX + 32 * 0.072 + 0.26, 6.18, 2.60, 0.22, "← 역전      개선 →", size=8,
   color=MUTED, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------- 오른쪽 35%: 프로파일 교차
RX, RW = 8.92, 3.91
tb(s, RX, 2.34, RW, 0.26, "RNNoise · U85 · TOTAL cycles", size=11, bold=True, color=NAVY)
bx, by, bw, bh = RX, 2.66, RW, 3.34
rect(s, bx, by, bw, bh, PALE)
X0, X1 = bx + 0.86, bx + bw - 0.70
YTOP, YBOT = by + 0.46, by + bh - 0.66
VMAX, VMIN = 66000.0, 28000.0


def cy(v):
    return YTOP + (YBOT - YTOP) * (VMAX - v) / (VMAX - VMIN)


for g in (60000, 50000, 40000, 30000):
    conn(s, bx + 0.70, cy(g), bx + bw - 0.14, cy(g), LINEC, 0.75, "sysDot")
    tb(s, bx + 0.08, cy(g) - 0.10, 0.58, 0.20, "%dk" % (g // 1000), size=8, color=MUTED,
       align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

L256, L512 = 36086, 31086
M256, M512 = 62086, 55086
conn(s, X0, cy(L256), X1, cy(L512), GREEN, 2.6)          # low 고정
conn(s, X0, cy(M256), X1, cy(M512), STEEL, 2.6)          # mid 고정
conn(s, X0, cy(L256), X1, cy(M512), RED, 2.8, "dash")    # 기본 조합
for x, v, colr, ha in ((X0, L256, GREEN, "l"), (X1, L512, GREEN, "r"),
                       (X0, M256, STEEL, "l"), (X1, M512, STEEL, "r")):
    dot(s, x, cy(v), 0.15, colr, WHITE)
tb(s, X0 - 0.86, cy(M256) - 0.12, 0.76, 0.24, "62.1k", size=9.5, bold=True, color=STEEL,
   align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
tb(s, X0 - 0.86, cy(L256) - 0.12, 0.76, 0.24, "36.1k", size=9.5, bold=True, color=GREEN,
   align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
tb(s, X1 + 0.10, cy(M512) - 0.12, 0.62, 0.24, "55.1k", size=9.5, bold=True, color=STEEL,
   anchor=MSO_ANCHOR.MIDDLE)
tb(s, X1 + 0.10, cy(L512) - 0.12, 0.62, 0.24, "31.1k", size=9.5, bold=True, color=GREEN,
   anchor=MSO_ANCHOR.MIDDLE)
tb(s, X0 - 0.60, YBOT + 0.10, 1.20, 0.24, "256 MAC", size=9, bold=True, color=GRAY,
   align=PP_ALIGN.CENTER)
tb(s, X1 - 0.60, YBOT + 0.10, 1.20, 0.24, "512 MAC", size=9, bold=True, color=GRAY,
   align=PP_ALIGN.CENTER)
MX = (X0 + X1) / 2
tb(s, MX - 0.95, (cy(M256) + cy(M512)) / 2 - 0.34, 1.90, 0.24, "mid 프로파일 고정",
   size=8.5, bold=True, color=STEEL, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
tb(s, MX - 0.95, (cy(L256) + cy(L512)) / 2 + 0.12, 1.90, 0.24, "low 프로파일 고정",
   size=8.5, bold=True, color=GREEN, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
tb(s, MX - 1.10, (cy(L256) + cy(M512)) / 2 - 0.36, 2.00, 0.24, "기본 조합 (실제 사다리)",
   size=8.5, bold=True, color=RED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print("wrote", OUT, "| 전이", len(TRANS), "| 역전 셀", len(REV))
