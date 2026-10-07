import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml.ns import qn

SRC, OUT = "user_v5.pptx", "Corstone_Family_v8.pptx"
NAVY = RGBColor(0x13, 0x36, 0x6B); TEXT = RGBColor(0x22, 0x36, 0x4B); GRAY = RGBColor(0x56, 0x65, 0x77)
LIGHT = RGBColor(0xEF, 0xF3, 0xF8); RED = RGBColor(0xA4, 0x0F, 0x16); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
KO, LAT = "NanumGothic", "Poppins"

prs = Presentation(SRC)
S = lambda n: prs.slides[n - 1]
LAYOUT = S(21).slide_layout


def _font(run, size, bold=False, color=TEXT):
    f = run.font; f.name = LAT; f.size = Pt(size); f.bold = bold; f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", KO)


def tb(slide, x, y, w, h, lines, size=15, bold=False, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.1):
    if isinstance(lines, str):
        lines = [lines]
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, item in enumerate(lines):
        text, opt = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        r = p.add_run(); r.text = text
        _font(r, opt.get("size", size), opt.get("bold", bold), opt.get("color", color))
    return box


def rect(slide, x, y, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill; s.line.fill.background(); s.shadow.inherit = False
    return s


def grid(slide, x, y, widths, rows, header_h=0.44, row_h=0.4, size=12.5, header_size=12.5, row_hs=None):
    cx = x
    for j, w in enumerate(widths):
        rect(slide, cx, y, w, header_h, NAVY)
        tb(slide, cx + 0.12, y + 0.06, w - 0.24, header_h - 0.1, rows[0][j], size=header_size, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        cx += w
    cy = y + header_h
    for i, row in enumerate(rows[1:]):
        h = row_hs[i] if row_hs else row_h
        cx = x
        for j, w in enumerate(widths):
            rect(slide, cx, cy, w, h, LIGHT if i % 2 == 0 else WHITE)
            tb(slide, cx + 0.12, cy + 0.05, w - 0.24, h - 0.1, row[j], size=size, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)
            cx += w
        cy += h
    return cy


def frame(slide, num, title, subtitle, takeaway):
    """Draw the v5 frame (title, badge, subtitle, takeaway bar) on a slide."""
    tb(slide, 0.48, 0.30, 11.70, 0.60, title, size=26, bold=True, color=WHITE)
    tb(slide, 12.03, 0.40, 0.80, 0.30, str(num), size=13, color=WHITE)
    tb(slide, 0.55, 1.34, 12.20, 0.57, subtitle, size=17, color=GRAY)
    rect(slide, 0.50, 6.35, 12.32, 0.50, LIGHT); rect(slide, 0.50, 6.35, 0.04, 0.50, NAVY)
    tb(slide, 0.65, 6.46, 12.00, 0.32, takeaway, size=14.5, bold=True, color=NAVY)


def clear(slide):
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)


def set_text(shape, lines):
    if isinstance(lines, str):
        lines = [lines]
    tf = shape.text_frame; p0 = tf.paragraphs[0]; r0 = p0.runs[0]
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    r0.text = lines[0]; rPr0 = r0._r.find(qn("a:rPr"))
    for extra in lines[1:]:
        p = tf.add_paragraph(); p.alignment = p0.alignment
        r = p.add_run(); r.text = extra
        if rPr0 is not None:
            r._r.insert(0, copy.deepcopy(rPr0))


def edit(n, mapping):
    shapes = list(S(n).shapes)
    for idx, lines in mapping.items():
        set_text(shapes[idx], lines)


