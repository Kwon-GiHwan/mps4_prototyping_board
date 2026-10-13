# -*- coding: utf-8 -*-
"""초안 -- 총비용 범위(optionA 왼쪽)와 비용 구성(v9 3장)을 한 장에.
두 차트는 같은 행 순서(총비용 내림차순)를 공유한다. 별도 파일이며
EthosU_RQ_v9.pptx 는 건드리지 않는다.
"""
import csv, json, math, os, tarfile
from collections import defaultdict

exec(open("v8_helpers.py", encoding="utf-8").read())
from pptx.enum.shapes import MSO_CONNECTOR

OUT = "EthosU_RQ_slide2_combined_build.pptx"
PAPER = "../../paper"
EXEC_TAR = PAPER + "/evidence/executability-20260825/cells/per_cell_evidence.tar.gz"
CYCLES_CSV = PAPER + "/raw_data_report/tables/3_3_cycles_by_mac.csv"
GROUPS_CSV = PAPER + "/mechanism/U85_GROUP_DIFFERENTIAL.csv"

PALE = RGBColor(0xF7, 0xF9, 0xFC)
LINEC = RGBColor(0xCB, 0xD6, 0xE3)
MUTED = RGBColor(0x8A, 0x97, 0xA6)
EYE = RGBColor(0xAE, 0xC2, 0xDE)
SPAN = RGBColor(0xC4, 0xCF, 0xDC)
SMALL = RGBColor(0x9F, 0xAE, 0xC0)

MODELS = [                                    # 총비용 내림차순 (두 차트 공통 행 순서)
    ("wav2letter_pruned_int8", "Wav2Letter", RGBColor(0x13, 0x36, 0x6B)),
    ("mobilenet_v2_1.0_224_INT8", "MobileNetV2", RGBColor(0x4A, 0x7A, 0xB8)),
    ("yolo-fastest_192_face_v4", "YOLO-Fastest", RGBColor(0x83, 0xAB, 0xD6)),
    ("ad_medium_int8", "AD MedNet", RGBColor(0x2E, 0x7D, 0x4F)),
    ("vww4_128_128_INT8", "VWW4", RGBColor(0x8C, 0x9A, 0xA8)),
    ("kws_micronet_m", "KWS MicroNet", RGBColor(0xC2, 0x8A, 0x1E)),
    ("rnnoise_INT8", "RNNoise", RGBColor(0xA4, 0x0F, 0x16)),
]
CFG = [("Corstone-300 · U55", "SSE-300", "ethos-u55"),
       ("Corstone-300 · U65", "SSE-300", "ethos-u65"),
       ("Corstone-310 · U55", "SSE-310", "ethos-u55"),
       ("Corstone-310 · U65", "SSE-310", "ethos-u65"),
       ("Corstone-315 · U65", "SSE-315", "ethos-u65"),
       ("Corstone-320 · U85", "SSE-320", "ethos-u85")]
CATS = [("Conv2D", "Conv2D", RGBColor(0x13, 0x36, 0x6B)),
        ("DepthwiseConv2D", "DepthwiseConv2D", RGBColor(0x4A, 0x7A, 0xB8)),
        ("혼합 그룹", "혼합 그룹", RGBColor(0xC2, 0x8A, 0x1E)),
        ("기타 단일 연산", "기타 단일 연산", RGBColor(0x8A, 0x97, 0xA6))]

# ------------------------------------------------------------------ 원자료
CYC = {}
_tf = tarfile.open(EXEC_TAR)
for _m in _tf.getmembers():
    _n = os.path.basename(_m.name)
    if not _n.endswith(".json") or _n.startswith("_"):
        continue
    _d = json.load(_tf.extractfile(_m))
    _wl, _plat, _nm = _d["cell_id"].split("__")
    _npu, _mac = _nm.rsplit("-", 1)
    for _a in _d["attempts"]:
        _r = _a.get("run") or {}
        if _r.get("status") == "SUCCESS" and _r.get("npu_total_cycles") is not None:
            CYC[(_plat, _npu, _wl, int(_mac))] = _r["npu_total_cycles"]
_k = 0
for _r in csv.DictReader(open(CYCLES_CSV, encoding="utf-8")):
    if _r["cycles"]:
        assert CYC[(_r["platform"], _r["npu"], _r["workload"], int(_r["mac"]))] == int(_r["cycles"])
        _k += 1
assert _k == 74

VAL = {m: [CYC[(p, u, m, 256)] for _l, p, u in CFG] for m, _n, _c in MODELS}
assert [m for m, _n, _c in MODELS] == sorted(VAL, key=lambda m: -max(VAL[m])), "행 순서"
_ref = None
for _l, _p, _u in CFG:
    _o = sorted(VAL, key=lambda w: CYC[(_p, _u, w, 256)])
    _ref = _o if _ref is None else _ref
    assert _o == _ref, _l


def _cat(t):
    st = set(t.split())
    st.discard("KERNEL_WAIT")
    if st == {"Conv2D"}:
        return "Conv2D"
    if st == {"DepthwiseConv2D"}:
        return "DepthwiseConv2D"
    return "혼합 그룹" if len(st) > 1 else "기타 단일 연산"


GRP = defaultdict(lambda: defaultdict(int))
GTOT = defaultdict(int)
for _r in csv.DictReader(open(GROUPS_CSV, encoding="utf-8")):
    if _r["binding_pair"] != "B-frozen":
        continue
    GRP[_r["workload"]][_cat(_r["member_types"])] += int(_r["cycles_256"])
    GTOT[_r["workload"]] += int(_r["cycles_256"])
COV = [GTOT[w] / CYC[("SSE-320", "ethos-u85", w, 256)] for w in GTOT]
assert 0.97 < min(COV) and max(COV) < 1.0
assert GRP["rnnoise_INT8"]["Conv2D"] == 0
COVTXT = "%.1f–%.1f%%" % (100 * min(COV), 100 * max(COV))


def blank():
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s


def conn(s, x1, y1, x2, y2, color=LINEC, w=0.75, dash=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                               Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": dash}))
    return c


def dot(s, cx, cy, d, color, ring=None):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2),
                           Inches(d), Inches(d))
    o.fill.solid(); o.fill.fore_color.rgb = color
    if ring is None:
        o.line.fill.background()
    else:
        o.line.color.rgb = ring; o.line.width = Pt(1.0)
    o.shadow.inherit = False
    return o


def tint(c, k=0.72):
    return RGBColor(*(int(c[i] + (255 - c[i]) * k) for i in range(3)))


def med(v):
    t = sorted(v)
    return (t[2] + t[3]) / 2


LO, HI = 4.13, 7.37


def cxl(v, x0, w):
    return x0 + w * (math.log10(v) - LO) / (HI - LO)


for _sid in list(prs.slides._sldIdLst):
    prs.part.drop_rel(_sid.rId)
    prs.slides._sldIdLst.remove(_sid)

s = blank()
tb(s, 0.48, 0.26, 8.90, 0.66, "2. 실행 비용은 얼마이고, 어디에 쓰이는가?", size=26,
   bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 9.45, 0.42, 3.40, 0.30, "RQ1 · 실행 비용", size=11.5, color=EYE, align=PP_ALIGN.RIGHT)
rect(s, 0.50, 1.24, 12.33, 0.84, LIGHT)
rect(s, 0.50, 1.24, 0.055, 0.84, NAVY)
tb(s, 0.74, 1.30, 11.90, 0.72,
   ["총비용 순위는 여섯 구성에서 모두 같지만, 그 비용이 어떤 연산 그룹에 쓰이는지는",
    "총비용 순위와 대응하지 않는다 — Conv 비중은 RNNoise 0%에서 KWS 83.8%까지 흩어진다."],
   size=15, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE, spacing=1.20)
tb(s, 0.52, 6.48, 12.30, 0.32,
   "왼쪽 256 MAC · TOTAL cycles · 여섯 구성의 최소–최대와 각 값, 굵은 점은 중앙값 · 공통 로그축   |   "
   "오른쪽 Corstone-320 · U85-256 · 계측 그룹 합계 대비 % (각 모델 TOTAL의 " + COVTXT + ")",
   size=9.5, color=MUTED)

ROW0, ROWH = 2.88, 0.48
LBX, LBW = 0.50, 1.54                      # 두 차트가 공유하는 모델 라벨
LIX, LIW = 2.26, 4.56                      # 왼쪽 range plot
BX, BW_ = 7.25, 5.58                       # 오른쪽 100% 누적 막대

tb(s, LBX, 2.24, LIX + LIW - LBX, 0.26, "총비용 범위 — 여섯 구성 · 256 MAC",
   size=11.5, bold=True, color=NAVY)
tb(s, BX, 2.24, BW_, 0.26, "비용 구성 — Corstone-320 · U85-256", size=11.5, bold=True,
   color=NAVY)

lx = BX
for _k2, nm, colr in CATS:
    rect(s, lx, 2.60, 0.18, 0.11, colr)
    tb(s, lx + 0.25, 2.51, 1.20, 0.26, nm, size=8.5, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)
    lx += 1.26

for g in (1e5, 1e6, 1e7):
    gx = cxl(g, LIX, LIW)
    conn(s, gx, ROW0, gx, ROW0 + 7 * ROWH, LINEC, 0.75, "sysDot")
    tb(s, gx - 0.40, ROW0 + 7 * ROWH + 0.03, 0.80, 0.22,
       {1e5: "100K", 1e6: "1M", 1e7: "10M"}[g], size=9, color=MUTED, align=PP_ALIGN.CENTER)

for i, (mid, name, colr) in enumerate(MODELS):
    yy = ROW0 + i * ROWH
    cy = yy + ROWH / 2
    tb(s, LBX, yy, LBW, ROWH, name, size=11, bold=True, color=colr,
       align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    v = VAL[mid]
    conn(s, cxl(min(v), LIX, LIW), cy, cxl(max(v), LIX, LIW), cy, SPAN, 3.4)
    for x in v:
        dot(s, cxl(x, LIX, LIW), cy, 0.09, SMALL, WHITE)
    dot(s, cxl(med(v), LIX, LIW), cy, 0.17, colr, WHITE)

    if mid in GTOT:
        cx = BX
        for key, _nm, ccol in CATS:
            f = GRP[mid][key] / GTOT[mid]
            if f <= 0:
                continue
            rect(s, cx, yy + 0.06, BW_ * f, ROWH - 0.14, ccol)
            if f >= 0.10:
                tb(s, cx, yy + 0.06, BW_ * f, ROWH - 0.14, "%.1f%%" % (100 * f), size=9,
                   bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            cx += BW_ * f

    else:
        rect(s, BX, yy + 0.06, BW_, ROWH - 0.14, RGBColor(0xF1, 0xF3, 0xF5))
        tb(s, BX, yy + 0.06, BW_, ROWH - 0.14, "그룹 프로파일 미계측", size=9, color=MUTED,
           align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print("wrote", OUT, "| 계측된 워크로드", len(GTOT))
