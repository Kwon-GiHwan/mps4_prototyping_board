# -*- coding: utf-8 -*-
"""3장 대안 -- 비용 구성을 '계측 단위(IRQ 서비스 창)' 기준으로 나눈 판.

현재 3장은 U85_GROUP_DIFFERENTIAL.csv 를 쓴다. 그 분할은 256<->512 차분을
정확히 내기 위한 공통 분할이라, 512 쪽 사정으로 합쳐진 것까지 256 그림에
반영된다. 이 장은 256 한 조건만 보여주므로 256_Low 의 계측 단위로 직접
나눈다. 남는 '혼합 유형'은 사이클이 측정되지 않은 것이 아니라, 한 계측 단위 안에
연산이 2개 이상 들어가 종류별 귀속만 불가능한 몫이다.

부수 효과: 이 분할은 256_Low 셀만 쓴다. P0-E2 에서 게이트가 거부한 셀은
512_Low / 512_Mid512 뿐이고 두 256_Low 셀은 PASS 이므로, 이 그림은 통과한
셀만으로 만들어진다.
"""
import csv, json, math, os, tarfile
from collections import defaultdict

exec(open("v8_helpers.py", encoding="utf-8").read())
from pptx.enum.shapes import MSO_CONNECTOR

OUT = "EthosU_RQ_slide3_units.pptx"
PAPER = "../../paper"
UNITS_CSV = PAPER + "/mechanism/U85_ATTRIBUTION_UNITS.csv"
CYCLES_CSV = PAPER + "/raw_data_report/tables/3_3_cycles_by_mac.csv"

MUTED = RGBColor(0x8A, 0x97, 0xA6)
EYE = RGBColor(0xAE, 0xC2, 0xDE)
AMBER = RGBColor(0xC2, 0x8A, 0x1E)

MODELS = [("wav2letter_pruned_int8", "Wav2Letter"),
          ("mobilenet_v2_1.0_224_INT8", "MobileNetV2"),
          ("yolo-fastest_192_face_v4", "YOLO-Fastest"),
          ("ad_medium_int8", "AD MedNet"),
          ("vww4_128_128_INT8", "VWW4"),
          ("kws_micronet_m", "KWS MicroNet"),
          ("rnnoise_INT8", "RNNoise")]
CATS = [("Conv2D", "Conv2D", RGBColor(0x13, 0x36, 0x6B)),
        ("DepthwiseConv2D", "DepthwiseConv2D", RGBColor(0x4A, 0x7A, 0xB8)),
        ("미분해", "혼합 유형 (2–6연산)", AMBER),
        ("비합성곱", "비합성곱", RGBColor(0x8A, 0x97, 0xA6))]


def cat(t):
    st = set(t.split()) - {"KERNEL_WAIT"}
    if st == {"Conv2D"}:
        return "Conv2D"
    if st == {"DepthwiseConv2D"}:
        return "DepthwiseConv2D"
    return "미분해" if len(st) > 1 else "비합성곱"


SHARE = defaultdict(lambda: defaultdict(int))
TOT = defaultdict(int)
MAXU = defaultdict(int)
NUNIT = defaultdict(int)
for r in csv.DictReader(open(UNITS_CSV, encoding="utf-8")):
    if r["binding"] != "256_Low":
        continue
    c, w = int(r["ccnt"]), r["workload"]
    SHARE[w][cat(r["op_types"])] += c
    TOT[w] += c
    MAXU[w] = max(MAXU[w], c)
    NUNIT[w] += 1
assert len(TOT) == 7, sorted(TOT)

WHOLE = {}
for r in csv.DictReader(open(CYCLES_CSV, encoding="utf-8")):
    if r["platform"] == "SSE-320" and r["mac"] == "256" and r["cycles"]:
        WHOLE[r["workload"]] = int(r["cycles"])
COV = [TOT[w] / WHOLE[w] for w in TOT]
assert 0.97 < min(COV) and max(COV) <= 1.0, COV
COVTXT = "%.1f–%.1f%%" % (100 * min(COV), 100 * max(COV))
assert SHARE["rnnoise_INT8"]["Conv2D"] == 0
NOTSPLIT = {w: SHARE[w]["미분해"] / TOT[w] for w in TOT}
OTHERMAX = max(v for w, v in NOTSPLIT.items() if w != "rnnoise_INT8")


def blank():
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s


for _sid in list(prs.slides._sldIdLst):
    prs.part.drop_rel(_sid.rId)
    prs.slides._sldIdLst.remove(_sid)

s = blank()
tb(s, 0.48, 0.26, 8.90, 0.66, "3. 실행 비용은 어떤 연산 그룹에 분포하는가?", size=26,
   bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 9.45, 0.42, 3.40, 0.30, "RQ1 · 비용 구성", size=11.5, color=EYE, align=PP_ALIGN.RIGHT)
rect(s, 0.50, 1.24, 12.33, 0.84, LIGHT)
rect(s, 0.50, 1.24, 0.055, 0.84, NAVY)
tb(s, 0.74, 1.30, 11.90, 0.72,
   ["Wav2Letter는 비용의 %.1f%%가 Conv이고, RNNoise는 Conv가 하나도 없다."
    % (100 * SHARE["wav2letter_pruned_int8"]["Conv2D"] / TOT["wav2letter_pruned_int8"]),
    "종류별로 가를 수 없는 몫은 RNNoise %.1f%%를 빼면 모두 %.0f%% 이하다."
    % (100 * NOTSPLIT["rnnoise_INT8"], math.ceil(100 * OTHERMAX))],
   size=15, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE, spacing=1.20)
tb(s, 0.52, 6.44, 12.30, 0.44,
   ["네 범주 모두 사이클은 측정됐다. 혼합 유형 = 한 계측 단위에 2–6개 연산이 함께 들어가 종류별 귀속만 안 되는 몫",
    "계측 단위(IRQ 서비스 창) 기준 · Corstone-320 · U85-256 의 256_Low 셀 단독 · "
    "계측 합계는 각 모델 TOTAL의 " + COVTXT + " · 출처 mechanism/U85_ATTRIBUTION_UNITS.csv"],
   size=9, color=MUTED, spacing=1.30)

lx = 2.36
for _k, nm, colr in CATS:
    rect(s, lx, 2.48, 0.20, 0.12, colr)
    tb(s, lx + 0.28, 2.40, 1.86, 0.26, nm, size=9.5, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)
    lx += 2.02
tb(s, 10.30, 2.40, 2.53, 0.26, "가장 큰 단일 단위 · 단위 수", size=9.5, bold=True,
   color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

BX, BW_ = 2.36, 7.76
NAME = {m: n for m, n in MODELS}
yy = 2.92
for w in sorted(TOT, key=lambda w: -SHARE[w]["Conv2D"] / TOT[w]):
    tb(s, 0.50, yy, 1.72, 0.44, NAME[w], size=11, bold=True, color=TEXT,
       align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    cx = BX
    for key, _nm, colr in CATS:
        f = SHARE[w][key] / TOT[w]
        if f <= 0:
            continue
        rect(s, cx, yy + 0.05, BW_ * f, 0.34, colr)
        if f >= 0.085:
            tb(s, cx, yy + 0.05, BW_ * f, 0.34, "%.1f%%" % (100 * f), size=9.5,
               bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        cx += BW_ * f
    tb(s, 10.30, yy, 2.53, 0.44, "%.1f%%   ·   %d개" % (100 * MAXU[w] / TOT[w], NUNIT[w]),
       size=10.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    yy += 0.47

prs.save(OUT)
print("wrote", OUT, "| 커버리지", COVTXT)
