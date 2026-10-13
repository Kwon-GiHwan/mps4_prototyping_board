# -*- coding: utf-8 -*-
"""초안 -- v9 4장 대안 2.
왼쪽 60% 모델별 인접 MAC 전이 speedup 점그래프 (점 하나 = 전이 하나),
오른쪽 40% RNNoise 프로파일 교차 (표기 최소화).
별도 파일이며 기존 덱은 건드리지 않는다.
"""
import csv, json, math, os, tarfile

exec(open("v8_helpers.py", encoding="utf-8").read())
from pptx.enum.shapes import MSO_CONNECTOR

OUT = "EthosU_RQ_slide4_dots.pptx"
PAPER = "../../paper"
EXEC_TAR = PAPER + "/evidence/executability-20260825/cells/per_cell_evidence.tar.gz"
CYCLES_CSV = PAPER + "/raw_data_report/tables/3_3_cycles_by_mac.csv"

MUTED = RGBColor(0x8A, 0x97, 0xA6)
EYE = RGBColor(0xAE, 0xC2, 0xDE)
LINEC = RGBColor(0xD2, 0xDB, 0xE5)
GREEN = RGBColor(0x2E, 0x7D, 0x4F)
STEEL = RGBColor(0x5B, 0x86, 0xBE)

MODELS = [("wav2letter_pruned_int8", "Wav2Letter"),
          ("mobilenet_v2_1.0_224_INT8", "MobileNetV2"),
          ("yolo-fastest_192_face_v4", "YOLO-Fastest"),
          ("ad_medium_int8", "AD MedNet"),
          ("vww4_128_128_INT8", "VWW4"),
          ("kws_micronet_m", "KWS MicroNet"),
          ("rnnoise_INT8", "RNNoise")]
HERO = "rnnoise_INT8"
GROUPS = [("SSE-300", "ethos-u55", [32, 64, 128, 256]),
          ("SSE-300", "ethos-u65", [256, 512]),
          ("SSE-310", "ethos-u55", [32, 64, 128, 256]),
          ("SSE-310", "ethos-u65", [256, 512]),
          ("SSE-315", "ethos-u65", [256, 512]),
          ("SSE-320", "ethos-u85", [128, 256, 512, 1024, 2048])]

CYC = {}
_tf = tarfile.open(EXEC_TAR)
for _m in _tf.getmembers():
    _n = os.path.basename(_m.name)
    if not _n.endswith(".json") or _n.startswith("_"):
        continue
    _d = json.load(_tf.extractfile(_m))
    _wl, _plat, _nm = _d["cell_id"].split("__")
    _npu, _mac = _nm.rsplit("-", 1)
    for _a in _d["attempts"]:
        _r = _a.get("run") or {}
        if _r.get("status") == "SUCCESS" and _r.get("npu_total_cycles") is not None:
            CYC[(_plat, _npu, _wl, int(_mac))] = _r["npu_total_cycles"]
_k = 0
for _r in csv.DictReader(open(CYCLES_CSV, encoding="utf-8")):
    if _r["cycles"]:
        assert CYC[(_r["platform"], _r["npu"], _r["workload"], int(_r["mac"]))] == int(_r["cycles"])
        _k += 1
assert _k == 74

TRANS = [(p, u, a, b) for p, u, macs in GROUPS for a, b in zip(macs, macs[1:])]
assert len(TRANS) == 13

SP = {}
for mid, _n in MODELS:
    vals = []
    for p, u, a, b in TRANS:
        x, y = CYC.get((p, u, mid, a)), CYC.get((p, u, mid, b))
        if x and y:
            vals.append(((p, u, a, b), x / y))
    SP[mid] = vals
REV = [(m, t, v) for m in SP for t, v in SP[m] if v < 1.0]
assert len(REV) == 2 and all(m == HERO for m, _t, _v in REV), REV
assert sorted(round(v, 2) for _m, _t, v in REV) == [0.66, 0.91]


def blank():
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s


def conn(s, x1, y1, x2, y2, color=LINEC, w=0.75, dash=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                               Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": dash}))
    return c


def dot(s, cx, cy_, d, color, ring=None):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy_ - d / 2),
                           Inches(d), Inches(d))
    o.fill.solid(); o.fill.fore_color.rgb = color
    if ring is None:
        o.line.fill.background()
    else:
        o.line.color.rgb = ring; o.line.width = Pt(0.9)
    o.shadow.inherit = False
    return o


def spread(xs, minsep):
    """값이 겹치면 같은 행 안에서 세로로만 조금 벌린다 (값은 바꾸지 않는다)."""
    order = sorted(range(len(xs)), key=lambda i: xs[i])
    off = [0.0] * len(xs)
    cluster = []
    for i in order:
        if cluster and xs[i] - xs[cluster[-1]] < minsep:
            cluster.append(i)
        else:
            if len(cluster) > 1:
                for k, j in enumerate(cluster):
                    off[j] = (k - (len(cluster) - 1) / 2.0)
            cluster = [i]
    if len(cluster) > 1:
        for k, j in enumerate(cluster):
            off[j] = (k - (len(cluster) - 1) / 2.0)
    return off


for _sid in list(prs.slides._sldIdLst):
    prs.part.drop_rel(_sid.rId)
    prs.slides._sldIdLst.remove(_sid)

s = blank()
tb(s, 0.48, 0.26, 8.90, 0.66, "4. MAC을 두 배 늘리면 얼마나 빨라지는가?", size=26,
   bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 9.45, 0.42, 3.40, 0.30, "RQ1 · MAC 확장", size=11.5, color=EYE, align=PP_ALIGN.RIGHT)
rect(s, 0.50, 1.24, 12.33, 0.84, LIGHT)
rect(s, 0.50, 1.24, 0.055, 0.84, NAVY)
tb(s, 0.74, 1.24, 11.90, 0.84,
   "RNNoise의 첫 역전은 같은 메모리 프로파일로 비교하면 나타나지 않는다.",
   size=15, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 0.52, 6.56, 12.30, 0.28,
   "점 하나 = 같은 플랫폼·NPU에서의 인접 MAC 전이 하나 (Corstone-300 · 310 · 315 · 320 전부)   |   "
   "오른쪽 단위 k cycles", size=9.5, color=MUTED)

# ------------------------------------------------ 왼쪽 60%: speedup 점그래프
LX, LBW = 0.50, 1.50
IX, IW = 2.14, 5.62
XLO, XHI = 0.45, 2.12
ROW0, ROWH = 2.62, 0.50


def cx(v):
    return IX + IW * (v - XLO) / (XHI - XLO)


tb(s, LX, 2.22, IX + IW - LX, 0.26, "모델별 MAC 2배 확장 이득", size=11.5, bold=True,
   color=NAVY)
for v, lab, colr, wdt, dsh in ((0.5, "0.5×", LINEC, 0.75, "sysDot"),
                               (1.0, "1×", NAVY, 1.4, None),
                               (1.5, "1.5×", LINEC, 0.75, "sysDot"),
                               (2.0, "2×", NAVY, 1.1, "dash")):
    conn(s, cx(v), ROW0 - 0.04, cx(v), ROW0 + 7 * ROWH + 0.02, colr, wdt, dsh)
    tb(s, cx(v) - 0.34, ROW0 + 7 * ROWH + 0.05, 0.68, 0.22, lab, size=9,
       bold=(v in (1.0, 2.0)), color=NAVY if v in (1.0, 2.0) else MUTED,
       align=PP_ALIGN.CENTER)
tb(s, cx(1.0) - 1.70, ROW0 + 7 * ROWH + 0.27, 1.60, 0.22, "← 역전", size=8.5,
   color=RED, align=PP_ALIGN.RIGHT)
tb(s, cx(1.0) + 0.10, ROW0 + 7 * ROWH + 0.27, 2.40, 0.22, "개선 →", size=8.5,
   color=NAVY)
tb(s, cx(2.0) - 2.62, ROW0 - 0.30, 2.60, 0.22, "2× = MAC 증가에 비례", size=8,
   color=MUTED, align=PP_ALIGN.RIGHT)

NOTE = []
for r, (mid, name) in enumerate(MODELS):
    yy = ROW0 + r * ROWH
    cyy = yy + ROWH / 2
    hero = mid == HERO
    tb(s, LX, yy, LBW, ROWH, name, size=11, bold=hero, color=RED if hero else TEXT,
       align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    vals = [v for _t, v in SP[mid]]
    offs = spread([cx(v) for v in vals], 0.10)
    for (t, v), o in zip(SP[mid], offs):
        rev = v < 1.0
        colr = RED if (hero and rev) else (RED if hero else STEEL)
        dot(s, cx(v), cyy + o * 0.098, 0.155 if (hero and rev) else 0.115, colr, WHITE)
        if hero and rev:
            NOTE.append("%.2f×  U85 %d→%d" % (v, t[2], t[3]))

NOTE.sort(reverse=True)
tb(s, 2.16, 4.90, 1.92, 0.66, [("역전 2건 — RNNoise", {"size": 8})] + NOTE, size=8.8,
   bold=True, color=RED, spacing=1.30)

# ------------------------------------------------ 오른쪽 40%: 프로파일 교차
RX, RW = 8.28, 4.55
tb(s, RX, 2.22, RW, 0.26, "RNNoise · U85 · 256 → 512", size=11.5, bold=True, color=NAVY)
X0, X1 = RX + 1.10, RX + RW - 0.92
YTOP, YBOT = 2.80, 5.86
VMAX, VMIN = 64000.0, 29000.0


def cyv(v):
    return YTOP + (YBOT - YTOP) * (VMAX - v) / (VMAX - VMIN)


L256, L512, M256, M512 = 36086, 31086, 62086, 55086
conn(s, X0, cyv(L256), X1, cyv(L512), GREEN, 1.6)
conn(s, X0, cyv(M256), X1, cyv(M512), STEEL, 1.6)
conn(s, X0, cyv(L256), X1, cyv(M512), RED, 3.2, "dash")
for x, v, colr in ((X0, L256, GREEN), (X1, L512, GREEN),
                   (X0, M256, STEEL), (X1, M512, STEEL)):
    dot(s, x, cyv(v), 0.135, colr, WHITE)
tb(s, X0 - 1.02, cyv(M256) - 0.13, 0.92, 0.26, "62.1", size=11, bold=True, color=STEEL,
   align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
tb(s, X0 - 1.02, cyv(L256) - 0.13, 0.92, 0.26, "36.1", size=11, bold=True, color=GREEN,
   align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
tb(s, X1 + 0.10, cyv(M512) - 0.13, 0.82, 0.26, "55.1", size=11, bold=True, color=STEEL,
   anchor=MSO_ANCHOR.MIDDLE)
tb(s, X1 + 0.10, cyv(L512) - 0.13, 0.82, 0.26, "31.1", size=11, bold=True, color=GREEN,
   anchor=MSO_ANCHOR.MIDDLE)
MXX = (X0 + X1) / 2
tb(s, MXX - 0.90, (cyv(M256) + cyv(M512)) / 2 - 0.32, 1.80, 0.24, "mid 고정", size=9.5,
   bold=True, color=STEEL, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
tb(s, MXX - 0.90, (cyv(L256) + cyv(L512)) / 2 + 0.10, 1.80, 0.24, "low 고정", size=9.5,
   bold=True, color=GREEN, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
tb(s, MXX - 1.34, (cyv(L256) + cyv(M512)) / 2 - 0.30, 2.00, 0.24, "기본 설정", size=10,
   bold=True, color=RED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
conn(s, X0 - 0.14, YBOT + 0.16, X1 + 0.14, YBOT + 0.16, LINEC, 0.9)
tb(s, X0 - 0.70, YBOT + 0.22, 1.40, 0.24, "256 MAC", size=9.5, bold=True, color=GRAY,
   align=PP_ALIGN.CENTER)
tb(s, X1 - 0.70, YBOT + 0.22, 1.40, 0.24, "512 MAC", size=9.5, bold=True, color=GRAY,
   align=PP_ALIGN.CENTER)

prs.save(OUT)
print("wrote", OUT, "| 점", sum(len(v) for v in SP.values()), "| 역전", len(REV))
