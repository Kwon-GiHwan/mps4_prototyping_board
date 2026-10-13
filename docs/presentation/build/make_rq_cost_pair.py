# -*- coding: utf-8 -*-
"""좌: 모델별 실행비용 범위(여섯 구성) · 우: 연산비용 구성(6계열 귀속).
두 차트가 같은 행 순서(총비용 내림차순)를 공유하고, 모델 이름은 왼쪽에 한 번만 둔다.
"""
import csv, json, math, os, tarfile
from collections import defaultdict

exec(open("v8_helpers.py", encoding="utf-8").read())
from pptx.enum.shapes import MSO_CONNECTOR

OUT = "EthosU_RQ_cost_pair.pptx"
PAPER = "../../paper"
EXEC_TAR = PAPER + "/evidence/executability-20260825/cells/per_cell_evidence.tar.gz"
CYCLES_CSV = PAPER + "/raw_data_report/tables/3_3_cycles_by_mac.csv"
UNITS_CSV = PAPER + "/mechanism/U85_ATTRIBUTION_UNITS.csv"

MUTED = RGBColor(0x8A, 0x97, 0xA6)
EYE = RGBColor(0xAE, 0xC2, 0xDE)
SPAN = RGBColor(0xC4, 0xCF, 0xDC)
SMALL = RGBColor(0x9F, 0xAE, 0xC0)

MODELS = [("wav2letter_pruned_int8", "Wav2Letter", RGBColor(0x13, 0x36, 0x6B)),
          ("mobilenet_v2_1.0_224_INT8", "MobileNetV2", RGBColor(0x4A, 0x7A, 0xB8)),
          ("yolo-fastest_192_face_v4", "YOLO-Fastest", RGBColor(0x83, 0xAB, 0xD6)),
          ("ad_medium_int8", "AD MedNet", RGBColor(0x2E, 0x7D, 0x4F)),
          ("vww4_128_128_INT8", "VWW4", RGBColor(0x8C, 0x9A, 0xA8)),
          ("kws_micronet_m", "KWS MicroNet", RGBColor(0xC2, 0x8A, 0x1E)),
          ("rnnoise_INT8", "RNNoise", RGBColor(0xA4, 0x0F, 0x16))]
CFG = [("SSE-300", "ethos-u55"), ("SSE-300", "ethos-u65"), ("SSE-310", "ethos-u55"),
       ("SSE-310", "ethos-u65"), ("SSE-315", "ethos-u65"), ("SSE-320", "ethos-u85")]
FAM = {"Conv2D": "Conv2D", "DepthwiseConv2D": "DWConv", "FullyConnected": "완전연결",
       "Add": "원소별", "Mul": "원소별", "Sub": "원소별",
       "Pad": "형상·이동", "Reshape": "형상·이동", "Pack": "형상·이동",
       "Concat": "형상·이동", "ResizeNearestNeighbor": "형상·이동",
       "AvgPool": "기타", "MaxPool": "기타", "Quantize": "기타", "Softmax": "기타"}
CATS = [("Conv2D", RGBColor(0x13, 0x36, 0x6B)),
        ("DWConv", RGBColor(0x4A, 0x7A, 0xB8)),
        ("완전연결", RGBColor(0x2E, 0x7D, 0x4F)),
        ("원소별", RGBColor(0x83, 0xAB, 0xD6)),
        ("형상·이동", RGBColor(0xC2, 0x8A, 0x1E)),
        ("기타", RGBColor(0x8A, 0x97, 0xA6))]
PRIO = [k for k, _c in CATS]

CYC = {}
_tf = tarfile.open(EXEC_TAR)
for _m in _tf.getmembers():
    _n = os.path.basename(_m.name)
    if not _n.endswith(".json") or _n.startswith("_"):
        continue
    _d = json.load(_tf.extractfile(_m))
    _wl, _p, _nm = _d["cell_id"].split("__")
    _npu, _mac = _nm.rsplit("-", 1)
    for _a in _d["attempts"]:
        _r = _a.get("run") or {}
        if _r.get("status") == "SUCCESS" and _r.get("npu_total_cycles") is not None:
            CYC[(_p, _npu, _wl, int(_mac))] = _r["npu_total_cycles"]
_k = 0
for _r in csv.DictReader(open(CYCLES_CSV, encoding="utf-8")):
    if _r["cycles"]:
        assert CYC[(_r["platform"], _r["npu"], _r["workload"], int(_r["mac"]))] == int(_r["cycles"])
        _k += 1
assert _k == 74
VAL = {m: [CYC[(p, u, m, 256)] for p, u in CFG] for m, _n, _c in MODELS}
assert [m for m, _n, _c in MODELS] == sorted(VAL, key=lambda m: -max(VAL[m])), "행 순서"

S = defaultdict(lambda: defaultdict(int))
T = defaultdict(int)
ASSUMED = defaultdict(int)
for r in csv.DictReader(open(UNITS_CSV, encoding="utf-8")):
    if r["binding"] != "256_Low":
        continue
    fams = {FAM[t] for t in r["op_types"].split() if t in FAM}
    c, w = int(r["ccnt"]), r["workload"]
    T[w] += c
    S[w][[f for f in PRIO if f in fams][0]] += c
    if len(fams) > 1:
        ASSUMED[w] += c
assert len(T) == 7
for w in T:
    assert sum(S[w][k] for k in PRIO) == T[w], w


def blank():
    s = prs.slides.add_slide(LAYOUT)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    return s


def conn(s, x1, y1, x2, y2, color, w=0.75, dash=None):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                               Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    if dash:
        ln = c.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": dash}))
    return c


def dot(s, cx, cy_, d, color, ring=None):
    o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy_ - d / 2),
                           Inches(d), Inches(d))
    o.fill.solid(); o.fill.fore_color.rgb = color
    if ring is None:
        o.line.fill.background()
    else:
        o.line.color.rgb = ring; o.line.width = Pt(0.9)
    o.shadow.inherit = False
    return o


def med(v):
    t = sorted(v)
    return (t[2] + t[3]) / 2


LO, HI = 4.13, 7.37
LBX, LBW = 0.50, 1.78
IX, IW = 2.36, 4.02
BX, BW_ = 7.02, 5.81
ROW0, ROWH = 2.86, 0.485


def cxl(v):
    return IX + IW * (math.log10(v) - LO) / (HI - LO)


for _sid in list(prs.slides._sldIdLst):
    prs.part.drop_rel(_sid.rId)
    prs.slides._sldIdLst.remove(_sid)

s = blank()
tb(s, 0.48, 0.26, 8.90, 0.66, "3. 실행 비용은 얼마이고, 어떤 연산에 쓰이는가?", size=26,
   bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
tb(s, 9.45, 0.42, 3.40, 0.30, "RQ1 · 실행 비용", size=11.5, color=EYE, align=PP_ALIGN.RIGHT)
rect(s, 0.50, 1.24, 12.33, 0.84, LIGHT)
rect(s, 0.50, 1.24, 0.055, 0.84, NAVY)
tb(s, 0.74, 1.30, 11.90, 0.72,
   ["모델 간 총비용 순위는 여섯 구성에서 모두 같다. 그러나 그 비용을 어떤 연산이 쓰는지는",
    "총비용 순위와 대응하지 않는다 — RNNoise만 합성곱이 없고 완전연결이 그 자리를 차지한다."],
   size=15, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE, spacing=1.20)
tb(s, 0.52, 6.56, 12.30, 0.30,
   "왼쪽 256 MAC · TOTAL cycles · 여섯 구성의 최소–최대와 각 값, 굵은 점은 중앙값 · 공통 로그축   |   "
   "오른쪽 Corstone-320 · U85-256 · 계측 단위 기준, 두 계열이 한 단위에 들면 무거운 쪽으로 귀속",
   size=9.5, color=MUTED)

tb(s, LBX, 2.22, IX + IW - LBX, 0.28, "모델 별 실행비용 (Cycles)", size=13, bold=True,
   color=NAVY)
tb(s, BX, 2.22, BW_, 0.28, "연산비용 구성 (%)", size=13, bold=True, color=NAVY)

lx = BX
for nm, colr in CATS:
    rect(s, lx, 2.58, 0.18, 0.11, colr)
    tb(s, lx + 0.24, 2.50, 0.92, 0.26, nm, size=8.5, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)
    lx += 0.97

for g in (1e5, 1e6, 1e7):
    conn(s, cxl(g), ROW0, cxl(g), ROW0 + 7 * ROWH, RGBColor(0xD2, 0xDB, 0xE5), 0.75, "sysDot")
    tb(s, cxl(g) - 0.40, ROW0 + 7 * ROWH + 0.05, 0.80, 0.22,
       {1e5: "100K", 1e6: "1M", 1e7: "10M"}[g], size=9, color=MUTED, align=PP_ALIGN.CENTER)

for i, (mid, name, colr) in enumerate(MODELS):
    yy = ROW0 + i * ROWH
    cy = yy + ROWH / 2
    tb(s, LBX, yy, LBW, ROWH, name, size=12, bold=True, color=colr,
       align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    v = VAL[mid]
    conn(s, cxl(min(v)), cy, cxl(max(v)), cy, SPAN, 3.4)
    for x in v:
        dot(s, cxl(x), cy, 0.085, SMALL, WHITE)
    dot(s, cxl(med(v)), cy, 0.165, colr, WHITE)

    cx = BX
    for nm, ccol in CATS:
        f = S[mid][nm] / T[mid]
        if f <= 0:
            continue
        rect(s, cx, yy + 0.06, BW_ * f, 0.34, ccol)
        if f >= 0.085:
            tb(s, cx, yy + 0.06, BW_ * f, 0.34, "%.1f%%" % (100 * f), size=9.5,
               bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        cx += BW_ * f

prs.save(OUT)
print("wrote", OUT)
